#!/usr/bin/env python3
"""
chm2word — Batch convert Microsoft Compiled HTML Help (.chm) files to Word (.docx).

Designed for converting licensed/authorized CHM documentation (e.g. E3 FDC help)
into editable Word documents while preserving the CHM table-of-contents structure
as Word heading levels (so the document has a navigable outline).

Cross-platform (Windows / Linux / macOS). It auto-detects whichever backends are
available, so you do not need every dependency installed:

  CHM extraction backend (first available wins, or choose with --backend):
    1. hh        : Windows native `hh.exe -decompile`         (best fidelity, Windows only)
    2. 7z        : 7-Zip / p7zip   `7z x`                       (recommended, cross-platform)
    3. chmlib    : `extract_chmLib` from chmlib                 (Linux/macOS)
    4. pychm     : Python `chm` package (libchm bindings)       (optional)

  DOCX backend (choose with --docx-backend):
    - word       : Microsoft Word via COM automation (Windows + pywin32)
                   — highest fidelity, no pandoc needed
    - pandoc     : `pandoc` (high fidelity: images, tables, formatting)
    - python     : pure-Python fallback (python-docx + beautifulsoup4 + lxml)

Typical usage (Windows PowerShell / CMD):
    python chm2word.py "C:\\Help\\E3_FDC.chm" "C:\\out"
    python chm2word.py "C:\\Help"           "C:\\out"     # batch a whole folder

Run `python chm2word.py --help` for all options.
"""

from __future__ import annotations

import argparse
import html
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.parse
from dataclasses import dataclass, field
from pathlib import Path

log = logging.getLogger("chm2word")


# --------------------------------------------------------------------------- #
# Small helpers
# --------------------------------------------------------------------------- #

def _which(name: str) -> str | None:
    return shutil.which(name)


def decode_bytes(data: bytes) -> str:
    """Decode CHM HTML bytes. CHM content is frequently cp1252/gbk/shift-jis
    rather than utf-8, so try a sensible ladder before lossy latin-1."""
    m = re.search(rb'charset\s*=\s*["\']?\s*([\w\-]+)', data[:4096], re.I)
    candidates = []
    if m:
        try:
            candidates.append(m.group(1).decode("ascii", "ignore").strip())
        except Exception:
            pass
    candidates += ["utf-8", "cp1252", "gb18030", "big5", "shift_jis", "latin-1"]
    seen: set[str] = set()
    for enc in candidates:
        enc = enc.lower()
        if not enc or enc in seen:
            continue
        seen.add(enc)
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("latin-1", "replace")


# --------------------------------------------------------------------------- #
# CHM extraction
# --------------------------------------------------------------------------- #

class ExtractionError(RuntimeError):
    pass


def extract_chm(chm_path: Path, dest: Path, backend: str) -> None:
    """Extract a .chm into `dest`. `backend` is one of
    auto|hh|7z|chmlib|pychm."""
    dest.mkdir(parents=True, exist_ok=True)
    order = ["hh", "7z", "chmlib", "pychm"] if backend == "auto" else [backend]

    last_err: Exception | None = None
    for be in order:
        try:
            if be == "hh":
                if os.name != "nt" or not _which("hh.exe"):
                    raise ExtractionError("hh.exe only available on Windows")
                # hh.exe -decompile <dir> <chm>  (returns immediately, sync enough)
                subprocess.run(["hh.exe", "-decompile", str(dest), str(chm_path)],
                               check=True)
            elif be == "7z":
                exe = _which("7z") or _which("7za") or _which("7zr")
                if not exe:
                    raise ExtractionError("7z/7za not found on PATH")
                subprocess.run([exe, "x", "-y", f"-o{dest}", str(chm_path)],
                               check=True, stdout=subprocess.DEVNULL,
                               stderr=subprocess.PIPE)
            elif be == "chmlib":
                exe = _which("extract_chmLib") or _which("extract_chmlib")
                if not exe:
                    raise ExtractionError("extract_chmLib not found on PATH")
                subprocess.run([exe, str(chm_path), str(dest)], check=True,
                               stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
            elif be == "pychm":
                _extract_with_pychm(chm_path, dest)
            else:
                raise ExtractionError(f"unknown backend {be!r}")

            if any(dest.rglob("*")):
                log.info("extracted %s using backend '%s'", chm_path.name, be)
                return
            raise ExtractionError(f"backend '{be}' produced no files")
        except Exception as e:  # try the next backend in 'auto' mode
            last_err = e
            log.debug("extraction backend '%s' failed: %s", be, e)

    raise ExtractionError(
        f"could not extract {chm_path.name}: {last_err}. "
        f"Install 7-Zip (recommended) or run on Windows with hh.exe.")


def _extract_with_pychm(chm_path: Path, dest: Path) -> None:
    try:
        from chm import chmlib  # type: ignore
    except Exception as e:
        raise ExtractionError(f"python 'chm' package not importable: {e}")

    f = chmlib.chm_open(str(chm_path).encode())
    if not f:
        raise ExtractionError("chm_open failed")
    files: list[str] = []

    def _cb(_chm, ui, _ctx):
        path = ui.path.decode("latin-1")
        if not path.endswith("/"):
            files.append(path)
        return chmlib.CHM_ENUMERATOR_CONTINUE

    chmlib.chm_enumerate(f, chmlib.CHM_ENUMERATE_NORMAL, _cb, None)
    for path in files:
        res, ui = chmlib.chm_resolve_object(f, path.encode())
        if res != chmlib.CHM_RESOLVE_SUCCESS:
            continue
        data = chmlib.chm_retrieve_object(f, ui, 0, ui.length)[1]
        out = dest / path.lstrip("/")
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_bytes(data)
    chmlib.chm_close(f)


# --------------------------------------------------------------------------- #
# Table of contents (.hhc) parsing
# --------------------------------------------------------------------------- #

@dataclass
class TocEntry:
    level: int
    title: str
    local: str | None  # path inside the extracted CHM (no anchor)


def find_hhc(root: Path) -> Path | None:
    hhcs = sorted(root.rglob("*.hhc"))
    return hhcs[0] if hhcs else None


def parse_hhc(hhc_path: Path) -> list[TocEntry]:
    """Parse the sitemap-style .hhc. It is malformed HTML built from nested
    <UL>/<LI><OBJECT><param name="Name"><param name="Local">."""
    raw = decode_bytes(hhc_path.read_bytes())
    entries: list[TocEntry] = []
    level = 0
    # Tokenize only the structural tags / objects we care about.
    token_re = re.compile(
        r"<\s*ul\b|<\s*/\s*ul\s*>|<\s*object\b[^>]*>(?P<obj>.*?)<\s*/\s*object\s*>",
        re.I | re.S)
    param_re = re.compile(
        r'<\s*param\b[^>]*\bname\s*=\s*"(?P<n>[^"]*)"[^>]*\bvalue\s*=\s*"(?P<v>[^"]*)"',
        re.I)
    for tok in token_re.finditer(raw):
        s = tok.group(0).lower()
        if s.startswith("<ul"):
            level += 1
        elif s.startswith("</ul") or s.startswith("< /ul"):
            level = max(0, level - 1)
        else:
            params = {m.group("n").lower(): m.group("v")
                      for m in param_re.finditer(tok.group("obj"))}
            name = params.get("name")
            if name is None:
                continue
            title = html.unescape(name).strip()
            local = params.get("local")
            if local:
                local = html.unescape(local).split("#", 1)[0].strip()
                local = urllib.parse.unquote(local).replace("\\", "/").lstrip("/")
            entries.append(TocEntry(max(1, level), title, local or None))
    return entries


def fallback_toc(root: Path) -> list[TocEntry]:
    """No .hhc: list every html page in stable path order, flat."""
    pages = sorted(
        p for p in root.rglob("*")
        if p.suffix.lower() in (".htm", ".html"))
    return [TocEntry(1, p.stem, str(p.relative_to(root)).replace("\\", "/"))
            for p in pages]


# --------------------------------------------------------------------------- #
# HTML assembly (one ordered HTML doc with TOC headings + bookmarks)
# --------------------------------------------------------------------------- #

@dataclass
class Assembled:
    html_path: Path
    resource_root: Path
    title: str
    page_count: int = 0
    sections: list[str] = field(default_factory=list)


def _resolve(root: Path, local: str) -> Path | None:
    cand = (root / local)
    if cand.is_file():
        return cand
    # CHM paths are case-insensitive; extracted FS may not be.
    target = local.lower()
    for p in root.rglob("*"):
        if p.is_file() and str(p.relative_to(root)).replace("\\", "/").lower() == target:
            return p
    return None


_BODY_RE = re.compile(r"<\s*body\b[^>]*>(?P<b>.*?)<\s*/\s*body\s*>", re.I | re.S)
_SCRIPT_RE = re.compile(r"<\s*(script|style)\b.*?<\s*/\s*\1\s*>", re.I | re.S)


def assemble_html(root: Path, toc: list[TocEntry], title: str,
                  out_dir: Path) -> Assembled:
    """Backward-compatible wrapper: assemble all TOC pages into one file."""
    return assemble(root, toc, title, out_dir / "_assembled.html")


def assemble(root: Path, toc: list[TocEntry], title: str,
             out_html: Path) -> Assembled:
    """Concatenate TOC-referenced pages into a single HTML file at
    `out_html`. Each TOC entry becomes an <hN> heading so Word gets a real
    outline; image/src references are rewritten to absolute paths so the
    docx/markdown backends can resolve them."""
    out_html.parent.mkdir(parents=True, exist_ok=True)
    parts: list[str] = [
        "<!DOCTYPE html><html><head><meta charset='utf-8'>"
        f"<title>{html.escape(title)}</title></head><body>"]
    seen_files: set[Path] = set()
    page_count = 0
    sections: list[str] = []

    for e in toc:
        hlvl = min(6, max(1, e.level))
        parts.append(f"<h{hlvl}>{html.escape(e.title)}</h{hlvl}>")
        sections.append(("  " * (e.level - 1)) + e.title)
        if not e.local:
            continue
        fp = _resolve(root, e.local)
        if not fp:
            log.warning("TOC entry not found in CHM: %s", e.local)
            continue
        # The same file can be referenced by several TOC nodes; emit body once.
        if fp in seen_files:
            continue
        seen_files.add(fp)

        body_html = decode_bytes(fp.read_bytes())
        m = _BODY_RE.search(body_html)
        inner = m.group("b") if m else body_html
        inner = _SCRIPT_RE.sub("", inner)
        inner = _rewrite_assets(inner, fp.parent, root)
        parts.append(f"<div class='chm-page'>{inner}</div>")
        page_count += 1

    parts.append("</body></html>")
    out_html.write_text("\n".join(parts), encoding="utf-8")
    return Assembled(out_html, root, title, page_count, sections)


_ATTR_RE = re.compile(
    r"""\b(src|href)\s*=\s*("([^"]*)"|'([^']*)')""", re.I)


def _rewrite_assets(body: str, page_dir: Path, root: Path) -> str:
    """Make relative img/src and local href absolute file paths so the docx
    backend can resolve images; drop in-CHM navigation hrefs to plain text."""
    def repl(m: re.Match) -> str:
        attr = m.group(1)
        val = (m.group(3) if m.group(3) is not None else m.group(4)) or ""
        low = val.lower()
        if low.startswith(("http:", "https:", "mailto:", "javascript:",
                           "data:", "#")):
            return m.group(0)
        clean = urllib.parse.unquote(val.split("#", 1)[0]).replace("\\", "/")
        if not clean:
            return m.group(0)
        target = (page_dir / clean).resolve()
        if not target.exists():
            target = _resolve(root, clean) or target
        if attr.lower() == "src" and target.exists():
            return f'{attr}="{target.as_uri()}"'
        if attr.lower() == "href":
            # internal cross-references become non-links (single flat doc)
            return ""
        return m.group(0)

    return _ATTR_RE.sub(repl, body)


# --------------------------------------------------------------------------- #
# DOCX backends
# --------------------------------------------------------------------------- #

def to_docx(asm: Assembled, out_docx: Path, backend: str,
            add_toc: bool) -> None:
    if backend == "word":
        _docx_word(asm, out_docx)
        return
    if backend == "pandoc":
        if not _which("pandoc"):
            raise RuntimeError("pandoc backend requested but pandoc not on PATH")
        _docx_pandoc(asm, out_docx, add_toc)
        return
    if backend == "python":
        _docx_python(asm, out_docx)
        return
    # auto: try Word (Windows) -> pandoc -> python, falling back on failure
    errors: list[str] = []
    if _word_available():
        try:
            _docx_word(asm, out_docx)
            return
        except Exception as e:
            errors.append(f"word: {e}")
    if _which("pandoc"):
        try:
            _docx_pandoc(asm, out_docx, add_toc)
            return
        except Exception as e:
            errors.append(f"pandoc: {e}")
    try:
        _docx_python(asm, out_docx)
        return
    except Exception as e:
        errors.append(f"python: {e}")
    raise RuntimeError("all docx backends failed: " + "; ".join(errors))


_word_app = None  # cached Word.Application instance for batch reuse


def _word_available() -> bool:
    if os.name != "nt":
        return False
    try:
        import win32com.client  # type: ignore  # noqa: F401
        return True
    except Exception:
        return False


def _get_word_app():
    """Lazily start Word.Application once; reuse across the whole batch."""
    global _word_app
    if _word_app is not None:
        return _word_app
    try:
        import win32com.client  # type: ignore
    except Exception as e:
        raise RuntimeError(
            "word backend needs pywin32 on Windows: pip install pywin32 "
            f"({e})")
    app = win32com.client.DispatchEx("Word.Application")
    app.Visible = False
    app.DisplayAlerts = 0  # wdAlertsNone
    import atexit

    def _quit():
        try:
            app.Quit(SaveChanges=0)
        except Exception:
            pass
    atexit.register(_quit)
    _word_app = app
    return app


def _docx_word(asm: Assembled, out_docx: Path) -> None:
    """Drive Microsoft Word via COM: open the assembled HTML, save as .docx.
    Word natively renders the HTML (including images referenced via file://
    URIs by `assemble()`), giving the highest fidelity on Windows."""
    app = _get_word_app()
    src = str(asm.html_path.resolve())
    dst = str(out_docx.resolve())
    doc = app.Documents.Open(src, ConfirmConversions=False, ReadOnly=False,
                             AddToRecentFiles=False, Visible=False)
    try:
        doc.SaveAs2(dst, FileFormat=16)  # 16 = wdFormatDocumentDefault (.docx)
    finally:
        doc.Close(SaveChanges=0)
    log.info("wrote %s (word COM)", out_docx.name)


def _docx_pandoc(asm: Assembled, out_docx: Path, add_toc: bool) -> None:
    cmd = ["pandoc", str(asm.html_path), "-f", "html", "-t", "docx",
           "--resource-path", str(asm.resource_root),
           "--extract-media", str(asm.resource_root),
           "-o", str(out_docx)]
    if add_toc:
        cmd += ["--toc", "--toc-depth=3"]
    subprocess.run(cmd, check=True, stderr=subprocess.PIPE)
    log.info("wrote %s (pandoc)", out_docx.name)


def _docx_python(asm: Assembled, out_docx: Path) -> None:
    """Lower-fidelity pure-Python path. Maps headings, paragraphs, lists,
    tables, bold/italic and images to a .docx via python-docx."""
    try:
        from bs4 import BeautifulSoup  # type: ignore
        from docx import Document  # type: ignore
        from docx.shared import Pt
    except Exception as e:
        raise RuntimeError(
            "python docx backend needs: pip install python-docx beautifulsoup4 "
            f"lxml  ({e})")

    soup = BeautifulSoup(asm.html_path.read_text(encoding="utf-8"), "lxml")
    doc = Document()
    doc.add_heading(asm.title, 0)

    def add_runs(par, node):
        for child in node.children:
            name = getattr(child, "name", None)
            if name is None:
                txt = str(child)
                if txt.strip():
                    par.add_run(html.unescape(txt))
            elif name in ("b", "strong"):
                par.add_run(child.get_text()).bold = True
            elif name in ("i", "em"):
                par.add_run(child.get_text()).italic = True
            elif name == "br":
                par.add_run("\n")
            else:
                add_runs(par, child)

    body = soup.body or soup
    for el in body.find_all(
            ["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre",
             "table", "img"], recursive=True):
        try:
            if el.name in ("h1", "h2", "h3", "h4", "h5", "h6"):
                doc.add_heading(el.get_text(strip=True), int(el.name[1]))
            elif el.name == "li":
                add_runs(doc.add_paragraph(style="List Bullet"), el)
            elif el.name == "pre":
                r = doc.add_paragraph().add_run(el.get_text())
                r.font.name, r.font.size = "Consolas", Pt(9)
            elif el.name == "table":
                rows = el.find_all("tr")
                if not rows:
                    continue
                ncol = max(len(r.find_all(["td", "th"])) for r in rows)
                t = doc.add_table(rows=0, cols=ncol)
                t.style = "Table Grid"
                for r in rows:
                    cells = r.find_all(["td", "th"])
                    rc = t.add_row().cells
                    for i, c in enumerate(cells[:ncol]):
                        rc[i].text = c.get_text(strip=True)
            elif el.name == "img":
                src = (el.get("src") or "").replace("file://", "")
                src = urllib.parse.unquote(src)
                if src.startswith("/") or re.match(r"^[a-zA-Z]:", src):
                    p = Path(src)
                    if p.exists() and p.suffix.lower() in (
                            ".png", ".jpg", ".jpeg", ".gif", ".bmp"):
                        try:
                            doc.add_picture(str(p))
                        except Exception as ie:
                            log.debug("skip image %s: %s", p, ie)
            elif el.name == "p":
                if el.get_text(strip=True):
                    add_runs(doc.add_paragraph(), el)
        except Exception as ee:
            log.debug("element %s skipped: %s", el.name, ee)

    doc.save(str(out_docx))
    log.info("wrote %s (python fallback)", out_docx.name)


# --------------------------------------------------------------------------- #
# Orchestration
# --------------------------------------------------------------------------- #

def convert_one(chm: Path, out_dir: Path, args) -> Path:
    out_docx = out_dir / (chm.stem + ".docx")
    if out_docx.exists() and not args.overwrite:
        log.info("skip (exists, use --overwrite): %s", out_docx.name)
        return out_docx

    work_parent = Path(args.keep_html) if args.keep_html else None
    tmp = tempfile.mkdtemp(prefix="chm2word_")
    work = Path(tmp)
    try:
        extract_dir = work / "chm"
        extract_chm(chm, extract_dir, args.backend)

        if args.toc:
            hhc = find_hhc(extract_dir)
            toc = parse_hhc(hhc) if hhc else fallback_toc(extract_dir)
            if hhc:
                log.info("TOC: %d entries from %s", len(toc), hhc.name)
            else:
                log.info("no .hhc; using flat page order (%d pages)", len(toc))
        else:
            toc = fallback_toc(extract_dir)

        asm = assemble_html(extract_dir, toc, chm.stem, work)
        log.info("assembled %d pages", asm.page_count)
        to_docx(asm, out_docx, args.docx_backend, add_toc=args.toc)

        if work_parent:
            keep = work_parent / chm.stem
            if keep.exists():
                shutil.rmtree(keep, ignore_errors=True)
            shutil.copytree(work, keep)
            log.info("kept intermediate HTML at %s", keep)
        return out_docx
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def gather_chms(inp: Path, recurse: bool) -> list[Path]:
    if inp.is_file():
        return [inp]
    pattern = "**/*.chm" if recurse else "*.chm"
    return sorted(inp.glob(pattern))


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="chm2word",
        description="Batch convert .chm help files to .docx, preserving the "
                    "CHM table-of-contents as Word heading structure.")
    p.add_argument("input", type=Path,
                   help="a .chm file OR a directory containing .chm files")
    p.add_argument("output", type=Path,
                   help="output directory for the .docx file(s)")
    p.add_argument("--backend", choices=["auto", "hh", "7z", "chmlib", "pychm"],
                   default="auto", help="CHM extraction backend (default auto)")
    p.add_argument("--docx-backend",
                   choices=["auto", "word", "pandoc", "python"],
                   default="auto",
                   help="auto picks Word(Windows)/pandoc/python in that order")
    p.add_argument("--no-toc", dest="toc", action="store_false",
                   help="ignore .hhc; just convert pages in file order")
    p.add_argument("--no-recurse", dest="recurse", action="store_false",
                   help="when input is a folder, do not recurse into subfolders")
    p.add_argument("--overwrite", action="store_true",
                   help="overwrite existing .docx outputs")
    p.add_argument("--keep-html", metavar="DIR",
                   help="also keep the intermediate extracted/assembled HTML")
    p.add_argument("-v", "--verbose", action="store_true")
    return p


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(levelname)s: %(message)s")

    if not args.input.exists():
        log.error("input not found: %s", args.input)
        return 2
    args.output.mkdir(parents=True, exist_ok=True)

    chms = gather_chms(args.input, args.recurse)
    if not chms:
        log.error("no .chm files found at %s", args.input)
        return 2
    log.info("found %d CHM file(s)", len(chms))

    ok, failed = 0, []
    for chm in chms:
        log.info("--- converting: %s", chm.name)
        try:
            convert_one(chm, args.output, args)
            ok += 1
        except Exception as e:
            log.error("FAILED %s: %s", chm.name, e)
            failed.append(chm.name)

    log.info("done: %d succeeded, %d failed", ok, len(failed))
    for f in failed:
        log.info("  failed: %s", f)
    return 0 if not failed else 1


if __name__ == "__main__":
    sys.exit(main())
